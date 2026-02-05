from io import BytesIO
from unittest.mock import patch
from django.urls import reverse
from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase
from rest_framework import status
import openpyxl

from apps.deliverables.models import (
    Project,
    ProjectMembership,
    SpecComparison,
    SpecComparisonStatus,
    SpecConflict,
    DrawingFile,
    DrawingExtraction,
    DrawingExtractionStatus,
    DrawingPage,
    DrawingNoteSection,
    DrawingNote,
)
from apps.teams.models import Team

User = get_user_model()


class TestSpecConflictsExport(APITestCase):
    """Test GET /projects/{id}/spec-conflicts/export/ endpoint"""

    def setUp(self):
        self.user = User.objects.create_user('test@example.com', password='testpass123')
        self.team = Team.objects.create(name='Test Team', slug='test-team')
        self.project = Project.objects.create(
            name='Test Project',
            project_number='P-001',
            team=self.team,
            created_by=self.user
        )
        self.project_version = self.project.versions.first()
        ProjectMembership.objects.create(
            project=self.project,
            user=self.user,
            role='project_member'
        )

        # Create drawing structures
        self.drawing_file = DrawingFile.objects.create(
            project=self.project,
            project_version=self.project_version,
            file_name='test.pdf',
            file_s3_key='drawings/test.pdf',
            md5='abc123',
        )
        self.extraction = DrawingExtraction.objects.create(
            drawing_file=self.drawing_file,
            status=DrawingExtractionStatus.SUCCESS,
        )
        self.page = DrawingPage.objects.create(
            drawing_file=self.drawing_file,
            extraction=self.extraction,
            page_number=1,
            page_type='drawing',
            extraction_status='success',
            sheet_number='P-201',
            sheet_title='Plumbing Plan',
        )
        self.section = DrawingNoteSection.objects.create(page=self.page, header='NOTES')
        self.note = DrawingNote.objects.create(
            section=self.section,
            note_number=1,
            category='general',
            text='Test note'
        )

        self.comparison = SpecComparison.objects.create(
            project=self.project,
            project_version=self.project_version,
            status=SpecComparisonStatus.SUCCESS,
            event_id='test-event',
        )
        self.conflict = SpecConflict.objects.create(
            comparison=self.comparison,
            note=self.note,
            note_id_from_lambda='1',
            note_text='Ball valve required',
            spec_text='Gate valve specified',
            spec_file_s3_key='specs/plumbing.pdf',
            spec_page_number=15,
            spec_masterformat_number='220500',
            confidence=0.85,
            reason='Valve type mismatch',
        )

        self.client.force_authenticate(user=self.user)
        self.url = reverse('deliverables:spec-conflicts-export', kwargs={'project_id': self.project.id})

    def test_returns_xlsx_content_type(self):
        """Test export returns correct content type"""
        response = self.client.get(self.url, {'project_version_id': self.project_version.id})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    def test_returns_attachment_disposition(self):
        """Test export returns attachment disposition"""
        response = self.client.get(self.url, {'project_version_id': self.project_version.id})

        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('spec_conflicts.xlsx', response['Content-Disposition'])

    def test_xlsx_contains_headers(self):
        """Test exported file contains correct headers"""
        response = self.client.get(self.url, {'project_version_id': self.project_version.id})

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        worksheet = workbook.active

        expected_headers = [
            'Drawing #', 'Sheet Title', 'Drawing Content', 'Related Spec Content',
            'Spec Section', 'Spec Page', 'Reason', 'Confidence'
        ]
        actual_headers = [cell.value for cell in worksheet[1]]

        self.assertEqual(actual_headers, expected_headers)

    def test_xlsx_contains_data(self):
        """Test exported file contains conflict data"""
        response = self.client.get(self.url, {'project_version_id': self.project_version.id})

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        worksheet = workbook.active

        # Check data row (row 2)
        self.assertEqual(worksheet.cell(2, 1).value, 'P-201')  # Drawing #
        self.assertEqual(worksheet.cell(2, 2).value, 'Plumbing Plan')  # Sheet Title
        self.assertEqual(worksheet.cell(2, 3).value, 'Ball valve required')  # Drawing Content
        self.assertEqual(worksheet.cell(2, 4).value, 'Gate valve specified')  # Spec Content
        self.assertEqual(worksheet.cell(2, 5).value, '220500')  # Spec Section
        self.assertEqual(worksheet.cell(2, 6).value, 15)  # Spec Page
        self.assertEqual(worksheet.cell(2, 7).value, 'Valve type mismatch')  # Reason
        self.assertEqual(worksheet.cell(2, 8).value, '85%')  # Confidence

    def test_respects_filters(self):
        """Test export respects filter parameters"""
        # Create another conflict that should be filtered out
        SpecConflict.objects.create(
            comparison=self.comparison,
            note_id_from_lambda='2',
            note_text='Other note',
            spec_text='Other spec',
            spec_file_s3_key='specs/other.pdf',
            spec_page_number=1,
            spec_masterformat_number='230500',
            confidence=0.7,
            reason='Different reason',
        )

        response = self.client.get(self.url, {
            'project_version_id': self.project_version.id,
            'spec_masterformat_number': '220500'
        })

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        worksheet = workbook.active

        # Should only have header + 1 data row
        self.assertEqual(worksheet.max_row, 2)

    def test_missing_project_version_id_returns_400(self):
        """Test 400 when project_version_id is missing"""
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_unauthenticated_returns_403(self):
        """Test 403 for unauthenticated requests"""
        self.client.force_authenticate(user=None)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
