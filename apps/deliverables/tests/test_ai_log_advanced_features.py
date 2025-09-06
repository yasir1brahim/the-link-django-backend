import json
import pytest
from io import BytesIO
from openpyxl import load_workbook
from django.urls import reverse
from rest_framework.test import APITestCase
from rest_framework import status
from django.contrib.auth import get_user_model
from apps.teams.models import Team, Membership as TeamMembership
from apps.deliverables.models import (
    Project, ProjectMembership, ProjectVersion, AiGeneratedLog,
    ROLE_PROJECT_ADMIN, ROLE_PROJECT_MEMBER
)
from apps.deliverables.serializers.specgpt import AiGeneratedLogSerializer


class AiGeneratedLogAdvancedFeaturesTests(APITestCase):
    """Test filter, search, and export functionality for AI generated logs"""
    
    def setUp(self):
        self.User = get_user_model()
        
        # Create test user
        self.user = self.User.objects.create_user(
            username='testuser',
            email='test@example.com',
            password='password123'
        )
        
        # Create team
        self.team = Team.objects.create(name='Test Team', slug='test-team')
        
        # Add user to team
        TeamMembership.objects.create(
            user=self.user,
            team=self.team,
            role='member'
        )
        
        # Create project
        self.project = Project.objects.create(
            name='Test Project',
            team=self.team
        )
        self.project_version = ProjectVersion.objects.get(project=self.project)
        
        # Add user to project
        ProjectMembership.objects.create(
            project=self.project,
            user=self.user,
            role=ROLE_PROJECT_MEMBER
        )
        
        # Create test QA planner log with structured data
        self.qa_log_data = [
            {
                'Spec Section #': '01 1000',
                'Spec Section Name': 'General Requirements',
                'Paragraph Number': '1.1',
                'item_type': 'requirement',
                'Requirement Text': 'All work must comply with local codes',
                'Responsible Party': 'General Contractor',
                'When Due': 'Before construction'
            },
            {
                'Spec Section #': '01 1000',
                'Spec Section Name': 'General Requirements',
                'Paragraph Number': '1.2',
                'item_type': 'submittal',
                'Requirement Text': 'Submit material certificates',
                'Responsible Party': 'Subcontractor',
                'When Due': 'Prior to installation'
            },
            {
                'Spec Section #': '02 2000',
                'Spec Section Name': 'Earthwork',
                'Paragraph Number': '2.1',
                'item_type': 'requirement',
                'Requirement Text': 'Excavate to specified depths',
                'Responsible Party': 'Excavation Contractor',
                'When Due': 'During excavation phase'
            }
        ]
        
        self.qa_log = AiGeneratedLog.objects.create(
            project=self.project,
            project_version=self.project_version,
            log_type='qa_planner',
            log_status='SUCCESS',
            log_data=self.qa_log_data
        )
        
        # Create test inspection log with structured data
        self.inspection_log_data = [
            {
                'Spec Section #': '03 3000',
                'Spec Section Name': 'Concrete',
                'Inspection Type And Requirements': 'Visual inspection of formwork',
                'Inspection Frequency': 'Daily',
                'Responsible Party': 'Quality Control Inspector'
            },
            {
                'Spec Section #': '05 5000',
                'Spec Section Name': 'Metals',
                'Inspection Type And Requirements': 'Welding inspection',
                'Inspection Frequency': 'Per weld',
                'Responsible Party': 'Certified Welding Inspector'
            }
        ]
        
        self.inspection_log = AiGeneratedLog.objects.create(
            project=self.project,
            project_version=self.project_version,
            log_type='inspection_log',
            log_status='SUCCESS',
            log_data=self.inspection_log_data
        )

    def test_filter_by_spec_section(self):
        """Test filtering by spec section number"""
        self.client.force_authenticate(user=self.user)
        
        # Create serializer and test filter functionality
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test filtering by single spec section
        filter_params = {'Spec Section #': ['01 1000']}
        filtered_data = serializer.filter_structured_data(self.qa_log_data, filter_params)
        
        self.assertEqual(len(filtered_data), 2)
        for item in filtered_data:
            self.assertEqual(item['Spec Section #'], '01 1000')
    
    def test_filter_by_multiple_spec_sections(self):
        """Test filtering by multiple spec section numbers (OR logic)"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test filtering by multiple spec sections (OR logic)
        filter_params = {'Spec Section #': ['01 1000', '02 2000']}
        filtered_data = serializer.filter_structured_data(self.qa_log_data, filter_params)
        
        self.assertEqual(len(filtered_data), 3)  # All 3 items should match
        spec_sections = {item['Spec Section #'] for item in filtered_data}
        self.assertEqual(spec_sections, {'01 1000', '02 2000'})
    
    def test_filter_by_item_type(self):
        """Test filtering by item type"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test filtering by item type
        filter_params = {'item_type': ['requirement']}
        filtered_data = serializer.filter_structured_data(self.qa_log_data, filter_params)
        
        self.assertEqual(len(filtered_data), 2)
        for item in filtered_data:
            self.assertEqual(item['item_type'], 'requirement')
    
    def test_filter_multiple_columns_and_logic(self):
        """Test filtering by multiple columns (AND logic between columns)"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test AND logic between different columns
        filter_params = {
            'Spec Section #': ['01 1000'],
            'item_type': ['requirement']
        }
        filtered_data = serializer.filter_structured_data(self.qa_log_data, filter_params)
        
        self.assertEqual(len(filtered_data), 1)
        item = filtered_data[0]
        self.assertEqual(item['Spec Section #'], '01 1000')
        self.assertEqual(item['item_type'], 'requirement')
    
    def test_filter_no_matches(self):
        """Test filtering that returns no matches"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test filter with no matches
        filter_params = {'Spec Section #': ['99 9999']}
        filtered_data = serializer.filter_structured_data(self.qa_log_data, filter_params)
        
        self.assertEqual(len(filtered_data), 0)
    
    def test_search_functionality(self):
        """Test search across all fields"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test search for "contractor" (should match multiple fields)
        search_results = serializer.search_structured_data(self.qa_log_data, 'contractor')
        
        self.assertEqual(len(search_results), 3)  # Should match "General Contractor", "Subcontractor", and "Excavation Contractor"
        
        # Test search for specific requirement text
        search_results = serializer.search_structured_data(self.qa_log_data, 'certificates')
        
        self.assertEqual(len(search_results), 1)
        self.assertIn('certificates', search_results[0]['Requirement Text'])
    
    def test_search_case_insensitive(self):
        """Test that search is case insensitive"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test case insensitive search for "GENERAL" - should match both "General Requirements" and "General Contractor"
        search_results = serializer.search_structured_data(self.qa_log_data, 'GENERAL')
        
        self.assertEqual(len(search_results), 2)  # Should match both "General Requirements" and "General Contractor"
        
        # Check that we found the expected matches
        found_section_names = {item['Spec Section Name'] for item in search_results}
        found_responsible_parties = {item['Responsible Party'] for item in search_results}
        self.assertIn('General Requirements', found_section_names)
        self.assertIn('General Contractor', found_responsible_parties)
    
    def test_search_partial_match(self):
        """Test that search finds partial matches"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test partial match search
        search_results = serializer.search_structured_data(self.qa_log_data, 'earth')
        
        self.assertEqual(len(search_results), 1)
        self.assertEqual(search_results[0]['Spec Section Name'], 'Earthwork')
    
    def test_search_no_matches(self):
        """Test search with no matches"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test search with no matches
        search_results = serializer.search_structured_data(self.qa_log_data, 'nonexistent')
        
        self.assertEqual(len(search_results), 0)
    
    def test_sorting_functionality(self):
        """Test sorting by different fields"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Test sorting by spec section number (ascending)
        sorted_data = serializer.sort_structured_data(self.qa_log_data, 'Spec Section #', 'asc')
        
        self.assertEqual(len(sorted_data), 3)
        self.assertEqual(sorted_data[0]['Spec Section #'], '01 1000')
        self.assertEqual(sorted_data[2]['Spec Section #'], '02 2000')
        
        # Test sorting by spec section number (descending)
        sorted_data = serializer.sort_structured_data(self.qa_log_data, 'Spec Section #', 'desc')
        
        self.assertEqual(sorted_data[0]['Spec Section #'], '02 2000')
        self.assertEqual(sorted_data[2]['Spec Section #'], '01 1000')
    
    def test_combined_filter_search_sort(self):
        """Test combining filter, search, and sort operations"""
        serializer = AiGeneratedLogSerializer(self.qa_log, context={'request': None})
        
        # Start with all data
        data = self.qa_log_data
        
        # Apply filter (spec section 01 1000)
        filter_params = {'Spec Section #': ['01 1000']}
        data = serializer.filter_structured_data(data, filter_params)
        self.assertEqual(len(data), 2)
        
        # Apply search (looking for "comply" which should only match the first requirement)
        data = serializer.search_structured_data(data, 'comply')
        self.assertEqual(len(data), 1)
        
        # Apply sort (by paragraph number)
        data = serializer.sort_structured_data(data, 'Paragraph Number', 'asc')
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['Paragraph Number'], '1.1')
    
    def test_filter_values_endpoint(self):
        """Test the filter_values API endpoint"""
        self.client.force_authenticate(user=self.user)
        
        url = reverse('ai-generated-log-filter-values', kwargs={
            'project_id': self.project.id, 
            'pk': self.qa_log.id
        })
        response = self.client.get(url)
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Check that filter values are returned correctly
        self.assertIn('filter_values', response.data)
        filter_values = response.data['filter_values']
        
        # Check spec section values
        self.assertIn('Spec Section #', filter_values)
        spec_sections = filter_values['Spec Section #']
        self.assertIn('01 1000', spec_sections)
        self.assertIn('02 2000', spec_sections)
        
        # Check item type values
        self.assertIn('item_type', filter_values)
        item_types = filter_values['item_type']
        self.assertIn('requirement', item_types)
        self.assertIn('submittal', item_types)
    
    def test_filter_values_endpoint_unauthenticated(self):
        """Test that unauthenticated users cannot access filter values"""
        url = reverse('ai-generated-log-filter-values', kwargs={
            'project_id': self.project.id, 
            'pk': self.qa_log.id
        })
        response = self.client.get(url)
        
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
    
    def test_export_endpoint_unauthenticated(self):
        """Test that unauthenticated users cannot access export"""
        url = reverse('ai-generated-log-export', kwargs={
            'project_id': self.project.id, 
            'pk': self.qa_log.id
        })
        response = self.client.get(url)
        
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
    
    def test_export_endpoint_authenticated(self):
        """Test that authenticated users can export data"""
        self.client.force_authenticate(user=self.user)
        
        url = reverse('ai-generated-log-export', kwargs={
            'project_id': self.project.id, 
            'pk': self.qa_log.id
        })
        response = self.client.get(url)
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'], 
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('Qa Planner_Export.xlsx', response['Content-Disposition'])
    
    def test_export_with_filters(self):
        """Test export with filter parameters"""
        self.client.force_authenticate(user=self.user)
        
        url = reverse('ai-generated-log-export', kwargs={
            'project_id': self.project.id, 
            'pk': self.qa_log.id
        })
        
        # Export with spec section filter
        response = self.client.get(url, {
            'filter_spec_section__': '01 1000'  # Note: double underscore from frontend
        })
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Verify the Excel content
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook.active
        
        # Should have header row + 2 data rows (filtered results)
        self.assertEqual(len(list(worksheet.rows)), 3)  # 1 header + 2 data rows
    
    def test_export_with_search(self):
        """Test export with search parameter"""
        self.client.force_authenticate(user=self.user)
        
        url = reverse('ai-generated-log-export', kwargs={
            'project_id': self.project.id, 
            'pk': self.qa_log.id
        })
        
        # Export with search filter
        response = self.client.get(url, {
            'search': 'contractor'
        })
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Verify the Excel content
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook.active
        
        # Should have header row + 3 data rows (search results for "contractor")
        self.assertEqual(len(list(worksheet.rows)), 4)  # 1 header + 3 data rows
    
    def test_export_column_order_qa_planner(self):
        """Test that export maintains correct column order for QA planner"""
        self.client.force_authenticate(user=self.user)
        
        url = reverse('ai-generated-log-export', kwargs={
            'project_id': self.project.id, 
            'pk': self.qa_log.id
        })
        response = self.client.get(url)
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Verify the Excel column order
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook.active
        
        # Check header row column order
        expected_headers = [
            'Spec Section #', 'Spec Section Name', 'Paragraph Number',
            'item_type', 'Requirement Text', 'Responsible Party', 'When Due'
        ]
        
        actual_headers = [cell.value for cell in worksheet[1]]
        self.assertEqual(actual_headers, expected_headers)
    
    def test_export_column_order_inspection_log(self):
        """Test that export maintains correct column order for inspection log"""
        self.client.force_authenticate(user=self.user)
        
        url = reverse('ai-generated-log-export', kwargs={
            'project_id': self.project.id, 
            'pk': self.inspection_log.id
        })
        response = self.client.get(url)
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        
        # Verify the Excel column order
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook.active
        
        # Check header row column order
        expected_headers = [
            'Spec Section #', 'Spec Section Name', 'Inspection Type And Requirements',
            'Inspection Frequency', 'Responsible Party'
        ]
        
        actual_headers = [cell.value for cell in worksheet[1]]
        self.assertEqual(actual_headers, expected_headers)
    
    def test_export_empty_data(self):
        """Test export with no data"""
        # Create log with no data
        empty_log = AiGeneratedLog.objects.create(
            project=self.project,
            project_version=self.project_version,
            log_type='qa_planner',
            log_status='SUCCESS',
            log_data=None
        )
        
        self.client.force_authenticate(user=self.user)
        
        url = reverse('ai-generated-log-export', kwargs={
            'project_id': self.project.id, 
            'pk': empty_log.id
        })
        response = self.client.get(url)
        
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('No data available for export', str(response.data))
    
    def test_export_no_matching_filters(self):
        """Test export with filters that match no data"""
        self.client.force_authenticate(user=self.user)
        
        url = reverse('ai-generated-log-export', kwargs={
            'project_id': self.project.id, 
            'pk': self.qa_log.id
        })
        
        # Export with filter that matches nothing
        response = self.client.get(url, {
            'filter_spec_section__': '99 9999'
        })
        
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('No data matches the specified filters', str(response.data))
    
    def test_export_nonexistent_log(self):
        """Test export with nonexistent log ID"""
        self.client.force_authenticate(user=self.user)
        
        url = reverse('ai-generated-log-export', kwargs={
            'project_id': self.project.id, 
            'pk': 99999
        })
        response = self.client.get(url)
        
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertIn('Log not found', str(response.data))
    
    def test_get_valid_sort_fields(self):
        """Test that get_valid_sort_fields returns correct fields for each log type"""
        from apps.deliverables.views.specgpt_views import AiGeneratedLogViewSet
        
        # Test inspection_log fields
        viewset = AiGeneratedLogViewSet()
        viewset.kwargs = {'pk': self.inspection_log.id}
        fields = viewset.get_valid_sort_fields()
        expected_fields = [
            'created_at', 'spec_section_number', 'spec_section_name',
            'inspection_type_and_requirements', 'inspection_frequency', 'responsible_party'
        ]
        self.assertEqual(fields, expected_fields)
        
        # Test qa_planner fields
        viewset = AiGeneratedLogViewSet()
        viewset.kwargs = {'pk': self.qa_log.id}
        fields = viewset.get_valid_sort_fields()
        expected_fields = [
            'created_at', 'spec_section_number', 'spec_section_name',
            'paragraph_number', 'item_type', 'requirement_text', 'responsible_party', 'when_due'
        ]
        self.assertEqual(fields, expected_fields)
    
    def test_get_filterable_columns(self):
        """Test that get_filterable_columns returns correct columns for each log type"""
        from apps.deliverables.views.specgpt_views import AiGeneratedLogViewSet
        
        viewset = AiGeneratedLogViewSet()
        
        # Test qa_planner filterable columns
        columns = viewset.get_filterable_columns('qa_planner')
        expected_columns = ['Spec Section #', 'item_type', 'Responsible Party']
        self.assertEqual(columns, expected_columns)
        
        # Test inspection_log filterable columns
        columns = viewset.get_filterable_columns('inspection_log')
        expected_columns = ['Spec Section #', 'Responsible Party']
        self.assertEqual(columns, expected_columns)
        
        # Test owner_deliverables_log filterable columns
        columns = viewset.get_filterable_columns('owner_deliverables_log')
        expected_columns = ['Spec Section #', 'Responsible Party', 'Deliverable Type']
        self.assertEqual(columns, expected_columns)
