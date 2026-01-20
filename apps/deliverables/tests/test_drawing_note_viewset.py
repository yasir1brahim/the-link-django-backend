# apps/deliverables/tests/test_drawing_note_viewset.py
from django.test import TestCase
from django.urls import reverse
from rest_framework.test import APIClient
from rest_framework import status
from apps.deliverables.models import (
    Project, ProjectMembership, ROLE_PROJECT_MEMBER,
    DrawingFile, DrawingExtraction, DrawingPage, DrawingNoteSection, DrawingNote,
    DrawingExtractionStatus, DrawingPageType, DrawingPageExtractionStatus,
)
from apps.teams.models import Team
from django.contrib.auth import get_user_model

User = get_user_model()


class TestDrawingNoteViewSet(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user('test@example.com', password='testpass')
        self.other_user = User.objects.create_user('other@example.com', password='testpass')
        self.team = Team.objects.create(name="Test Team", slug="test-team")
        self.project = Project.objects.create(
            name="Test Project",
            project_number="P-001",
            team=self.team,
            created_by=self.user
        )
        self.project_version = self.project.versions.first()
        ProjectMembership.objects.create(
            project=self.project,
            user=self.user,
            role=ROLE_PROJECT_MEMBER
        )

        # Create drawing data
        self.drawing_file = DrawingFile.objects.create(
            project=self.project,
            project_version=self.project_version,
            file_name="Mechanical.pdf",
            file_s3_key="drawings/test.pdf",
            md5="abc123",
        )
        self.extraction = DrawingExtraction.objects.create(
            drawing_file=self.drawing_file,
            status=DrawingExtractionStatus.SUCCESS,
        )
        self.page = DrawingPage.objects.create(
            drawing_file=self.drawing_file,
            extraction=self.extraction,
            page_number=1,
            page_type=DrawingPageType.DRAWING,
            extraction_status=DrawingPageExtractionStatus.SUCCESS,
        )
        self.section = DrawingNoteSection.objects.create(
            page=self.page,
            header="GENERAL NOTES:",
        )
        self.note = DrawingNote.objects.create(
            section=self.section,
            note_number=1,
            category="GENERAL NOTES",
            text="Test note content",
        )

        self.client.force_authenticate(user=self.user)

    def test_list_notes(self):
        """Test listing drawing notes for a project"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('results', response.data)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['text'], "Test note content")

    def test_list_notes_unauthenticated_denied(self):
        """Test unauthenticated access is denied"""
        self.client.logout()
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        # DRF may return 401 or 403 depending on authentication backend
        self.assertIn(response.status_code, [status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN])

    def test_list_notes_non_member_denied(self):
        """Test non-member access is denied"""
        self.client.force_authenticate(user=self.other_user)
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_filter_by_project_version(self):
        """Test filtering by project_version_id"""
        # Create note in different version
        version2 = self.project.versions.create(version_number=2, version_name="V2")
        drawing_file2 = DrawingFile.objects.create(
            project=self.project,
            project_version=version2,
            file_name="Electrical.pdf",
            file_s3_key="drawings/test2.pdf",
            md5="def456",
        )
        extraction2 = DrawingExtraction.objects.create(
            drawing_file=drawing_file2,
            status=DrawingExtractionStatus.SUCCESS,
        )
        page2 = DrawingPage.objects.create(
            drawing_file=drawing_file2,
            extraction=extraction2,
            page_number=1,
            page_type=DrawingPageType.DRAWING,
            extraction_status=DrawingPageExtractionStatus.SUCCESS,
        )
        section2 = DrawingNoteSection.objects.create(page=page2, header="ELECTRICAL:")
        DrawingNote.objects.create(
            section=section2,
            note_number=1,
            category="ELECTRICAL",
            text="Electrical note",
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})

        # Filter by first version
        response = self.client.get(url, {'project_version_id': self.project_version.id})
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['category'], "GENERAL NOTES")

        # Filter by second version
        response = self.client.get(url, {'project_version_id': version2.id})
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['category'], "ELECTRICAL")

    def test_filter_by_category(self):
        """Test filtering by category"""
        # Create note with different category
        DrawingNote.objects.create(
            section=self.section,
            note_number=2,
            category="MECHANICAL",
            text="Mechanical note",
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'category': 'MECHANICAL'})

        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['category'], "MECHANICAL")

    def test_all_filter_vals_returned(self):
        """Test that filter options are returned in response"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertIn('all_filter_vals', response.data)
        self.assertIn('category', response.data['all_filter_vals'])
        self.assertIn('drawing_files', response.data['all_filter_vals'])

        # drawing_files should be array of {id, name} objects
        drawing_files = response.data['all_filter_vals']['drawing_files']
        self.assertTrue(len(drawing_files) > 0)
        self.assertIn('id', drawing_files[0])
        self.assertIn('name', drawing_files[0])
        self.assertEqual(drawing_files[0]['id'], self.drawing_file.id)
        self.assertEqual(drawing_files[0]['name'], "Mechanical.pdf")

    def test_processing_status_returned(self):
        """Test that processing_status is included in response"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertIn('processing_status', response.data)
        status_data = response.data['processing_status']
        self.assertIn('is_processing', status_data)
        self.assertIn('files_processing', status_data)
        self.assertIn('files_completed', status_data)
        self.assertIn('files_failed', status_data)
        self.assertIn('files', status_data)

    def test_processing_status_shows_processing_files(self):
        """Test processing_status correctly shows files being processed"""
        # Create a file that is still processing
        processing_file = DrawingFile.objects.create(
            project=self.project,
            project_version=self.project_version,
            file_name="Electrical.pdf",
            file_s3_key="drawings/electrical.pdf",
            md5="def456",
        )
        DrawingExtraction.objects.create(
            drawing_file=processing_file,
            status=DrawingExtractionStatus.PROCESSING,
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        status_data = response.data['processing_status']
        self.assertTrue(status_data['is_processing'])
        self.assertEqual(status_data['files_processing'], 1)
        self.assertEqual(status_data['files_completed'], 1)  # Original file
        self.assertEqual(len(status_data['files']), 1)
        self.assertEqual(status_data['files'][0]['name'], "Electrical.pdf")
        self.assertEqual(status_data['files'][0]['status'], "PROCESSING")

    def test_export_notes_to_xlsx(self):
        """Test exporting drawing notes to XLSX format"""
        import openpyxl
        from io import BytesIO

        url = reverse('deliverables:drawing-note-export', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment; filename=drawing_notes.xlsx', response['Content-Disposition'])

        # Verify the workbook content
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        worksheet = workbook.active

        # Check headers
        headers = ['File Name', 'Sheet #', 'Sheet Title', 'Page', 'Section', 'Note #', 'Category', 'Note Text']
        for col, header in enumerate(headers, start=1):
            self.assertEqual(worksheet.cell(row=1, column=col).value, header)

        # Check data row
        self.assertEqual(worksheet.cell(row=2, column=1).value, "Mechanical.pdf")  # File Name
        # Sheet # and Sheet Title are empty strings when null (openpyxl may interpret '' as None)
        self.assertIn(worksheet.cell(row=2, column=2).value, ["", None])  # Sheet #
        self.assertIn(worksheet.cell(row=2, column=3).value, ["", None])  # Sheet Title
        self.assertEqual(worksheet.cell(row=2, column=4).value, 1)  # Page
        self.assertEqual(worksheet.cell(row=2, column=5).value, "GENERAL NOTES:")  # Section
        self.assertEqual(worksheet.cell(row=2, column=6).value, 1)  # Note #
        self.assertEqual(worksheet.cell(row=2, column=7).value, "GENERAL NOTES")  # Category
        self.assertEqual(worksheet.cell(row=2, column=8).value, "Test note content")  # Note Text

    def test_export_notes_with_filter(self):
        """Test exporting filtered drawing notes"""
        import openpyxl
        from io import BytesIO

        # Create additional note with different category
        DrawingNote.objects.create(
            section=self.section,
            note_number=2,
            category="MECHANICAL",
            text="Mechanical note content",
        )

        url = reverse('deliverables:drawing-note-export', kwargs={'project_id': self.project.id})

        # Export with category filter
        response = self.client.get(url, {'category': 'MECHANICAL'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        worksheet = workbook.active

        # Should only have header row + 1 data row (filtered result)
        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual(worksheet.cell(row=2, column=7).value, "MECHANICAL")  # Category is now column 7

    def test_export_notes_unauthenticated_denied(self):
        """Test export is denied for unauthenticated users"""
        self.client.logout()
        url = reverse('deliverables:drawing-note-export', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertIn(response.status_code, [status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN])

    def test_export_notes_non_member_denied(self):
        """Test export is denied for non-members"""
        self.client.force_authenticate(user=self.other_user)
        url = reverse('deliverables:drawing-note-export', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_sort_by_category_asc(self):
        """Test sorting by category ascending"""
        # Create notes with different categories
        DrawingNote.objects.create(
            section=self.section,
            note_number=2,
            category="MECHANICAL",
            text="Mechanical note",
        )
        DrawingNote.objects.create(
            section=self.section,
            note_number=3,
            category="ELECTRICAL",
            text="Electrical note",
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'category', 'sort_direction': 'asc'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        self.assertEqual(len(results), 3)
        # Alphabetical order: ELECTRICAL, GENERAL NOTES, MECHANICAL
        self.assertEqual(results[0]['category'], 'ELECTRICAL')
        self.assertEqual(results[1]['category'], 'GENERAL NOTES')
        self.assertEqual(results[2]['category'], 'MECHANICAL')

    def test_sort_by_category_desc(self):
        """Test sorting by category descending"""
        # Create notes with different categories
        DrawingNote.objects.create(
            section=self.section,
            note_number=2,
            category="MECHANICAL",
            text="Mechanical note",
        )
        DrawingNote.objects.create(
            section=self.section,
            note_number=3,
            category="ELECTRICAL",
            text="Electrical note",
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'category', 'sort_direction': 'desc'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        self.assertEqual(len(results), 3)
        # Reverse alphabetical: MECHANICAL, GENERAL NOTES, ELECTRICAL
        self.assertEqual(results[0]['category'], 'MECHANICAL')
        self.assertEqual(results[1]['category'], 'GENERAL NOTES')
        self.assertEqual(results[2]['category'], 'ELECTRICAL')

    def test_sort_by_text_asc(self):
        """Test sorting by text ascending"""
        DrawingNote.objects.create(
            section=self.section,
            note_number=2,
            category="GENERAL NOTES",
            text="Alpha note",
        )
        DrawingNote.objects.create(
            section=self.section,
            note_number=3,
            category="GENERAL NOTES",
            text="Zebra note",
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'text', 'sort_direction': 'asc'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        self.assertEqual(len(results), 3)
        self.assertEqual(results[0]['text'], 'Alpha note')
        self.assertEqual(results[1]['text'], 'Test note content')
        self.assertEqual(results[2]['text'], 'Zebra note')

    def test_sort_by_drawing_file_name_asc(self):
        """Test sorting by drawing file name ascending"""
        # Create another drawing file with notes
        drawing_file2 = DrawingFile.objects.create(
            project=self.project,
            project_version=self.project_version,
            file_name="Architectural.pdf",
            file_s3_key="drawings/arch.pdf",
            md5="xyz789",
        )
        extraction2 = DrawingExtraction.objects.create(
            drawing_file=drawing_file2,
            status=DrawingExtractionStatus.SUCCESS,
        )
        page2 = DrawingPage.objects.create(
            drawing_file=drawing_file2,
            extraction=extraction2,
            page_number=1,
            page_type=DrawingPageType.DRAWING,
            extraction_status=DrawingPageExtractionStatus.SUCCESS,
        )
        section2 = DrawingNoteSection.objects.create(page=page2, header="ARCH NOTES:")
        DrawingNote.objects.create(
            section=section2,
            note_number=1,
            category="ARCHITECTURAL",
            text="Architectural note",
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'drawing_file_name', 'sort_direction': 'asc'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        self.assertEqual(len(results), 2)
        # Alphabetical: Architectural.pdf before Mechanical.pdf
        self.assertEqual(results[0]['drawing_file_name'], 'Architectural.pdf')
        self.assertEqual(results[1]['drawing_file_name'], 'Mechanical.pdf')

    def test_sort_by_drawing_file_name_desc(self):
        """Test sorting by drawing file name descending"""
        # Create another drawing file with notes
        drawing_file2 = DrawingFile.objects.create(
            project=self.project,
            project_version=self.project_version,
            file_name="Architectural.pdf",
            file_s3_key="drawings/arch.pdf",
            md5="xyz789",
        )
        extraction2 = DrawingExtraction.objects.create(
            drawing_file=drawing_file2,
            status=DrawingExtractionStatus.SUCCESS,
        )
        page2 = DrawingPage.objects.create(
            drawing_file=drawing_file2,
            extraction=extraction2,
            page_number=1,
            page_type=DrawingPageType.DRAWING,
            extraction_status=DrawingPageExtractionStatus.SUCCESS,
        )
        section2 = DrawingNoteSection.objects.create(page=page2, header="ARCH NOTES:")
        DrawingNote.objects.create(
            section=section2,
            note_number=1,
            category="ARCHITECTURAL",
            text="Architectural note",
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'drawing_file_name', 'sort_direction': 'desc'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        self.assertEqual(len(results), 2)
        # Reverse alphabetical: Mechanical.pdf before Architectural.pdf
        self.assertEqual(results[0]['drawing_file_name'], 'Mechanical.pdf')
        self.assertEqual(results[1]['drawing_file_name'], 'Architectural.pdf')

    def test_sort_default_direction_is_asc(self):
        """Test that default sort direction is ascending when not specified"""
        DrawingNote.objects.create(
            section=self.section,
            note_number=2,
            category="MECHANICAL",
            text="Mechanical note",
        )
        DrawingNote.objects.create(
            section=self.section,
            note_number=3,
            category="ELECTRICAL",
            text="Electrical note",
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'category'})  # No sort_direction

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        # Should be ascending by default
        self.assertEqual(results[0]['category'], 'ELECTRICAL')
        self.assertEqual(results[1]['category'], 'GENERAL NOTES')
        self.assertEqual(results[2]['category'], 'MECHANICAL')

    def test_sort_invalid_column_uses_default_order(self):
        """Test that invalid sort_column falls back to default ordering"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'invalid_column'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        # Should not raise error, just use default ordering

    def test_sort_with_filter(self):
        """Test that sorting works combined with filtering"""
        # Create notes with different categories
        DrawingNote.objects.create(
            section=self.section,
            note_number=2,
            category="GENERAL NOTES",
            text="Zebra note",
        )
        DrawingNote.objects.create(
            section=self.section,
            note_number=3,
            category="GENERAL NOTES",
            text="Alpha note",
        )
        DrawingNote.objects.create(
            section=self.section,
            note_number=4,
            category="MECHANICAL",
            text="Mechanical note",
        )

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {
            'category': 'GENERAL NOTES',
            'sort_column': 'text',
            'sort_direction': 'asc'
        })

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        # Should only include GENERAL NOTES, sorted by text
        self.assertEqual(len(results), 3)
        self.assertEqual(results[0]['text'], 'Alpha note')
        self.assertEqual(results[1]['text'], 'Test note content')
        self.assertEqual(results[2]['text'], 'Zebra note')


class TestDrawingNoteViewSetSheetFields(TestCase):
    """Tests for sheet_number and sheet_title filtering, sorting, and filter values"""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user('test@example.com', password='testpass')
        self.team = Team.objects.create(name="Test Team", slug="test-team")
        self.project = Project.objects.create(
            name="Test Project",
            project_number="P-001",
            team=self.team,
            created_by=self.user
        )
        self.project_version = self.project.versions.first()
        ProjectMembership.objects.create(
            project=self.project,
            user=self.user,
            role=ROLE_PROJECT_MEMBER
        )

        # Create drawing file and extraction
        self.drawing_file = DrawingFile.objects.create(
            project=self.project,
            project_version=self.project_version,
            file_name="Mechanical.pdf",
            file_s3_key="drawings/test.pdf",
            md5="abc123",
        )
        self.extraction = DrawingExtraction.objects.create(
            drawing_file=self.drawing_file,
            status=DrawingExtractionStatus.SUCCESS,
        )

        # Create pages with different sheet values
        self.page1 = DrawingPage.objects.create(
            drawing_file=self.drawing_file,
            extraction=self.extraction,
            page_number=1,
            page_type=DrawingPageType.DRAWING,
            extraction_status=DrawingPageExtractionStatus.SUCCESS,
            sheet_number="A-101",
            sheet_title="First Floor Plan",
        )
        self.page2 = DrawingPage.objects.create(
            drawing_file=self.drawing_file,
            extraction=self.extraction,
            page_number=2,
            page_type=DrawingPageType.DRAWING,
            extraction_status=DrawingPageExtractionStatus.SUCCESS,
            sheet_number="M-201",
            sheet_title="Mechanical Plan",
        )
        self.page3 = DrawingPage.objects.create(
            drawing_file=self.drawing_file,
            extraction=self.extraction,
            page_number=3,
            page_type=DrawingPageType.DRAWING,
            extraction_status=DrawingPageExtractionStatus.SUCCESS,
            sheet_number=None,  # No sheet number
            sheet_title=None,   # No sheet title
        )

        # Create sections and notes
        self.section1 = DrawingNoteSection.objects.create(page=self.page1, header="GENERAL NOTES:")
        self.section2 = DrawingNoteSection.objects.create(page=self.page2, header="MECHANICAL NOTES:")
        self.section3 = DrawingNoteSection.objects.create(page=self.page3, header="OTHER NOTES:")

        self.note1 = DrawingNote.objects.create(
            section=self.section1, note_number=1, category="GENERAL", text="Note on A-101"
        )
        self.note2 = DrawingNote.objects.create(
            section=self.section2, note_number=1, category="MECHANICAL", text="Note on M-201"
        )
        self.note3 = DrawingNote.objects.create(
            section=self.section3, note_number=1, category="OTHER", text="Note without sheet"
        )

        self.client.force_authenticate(user=self.user)

    def test_filter_by_sheet_number_exact(self):
        """Test filtering by sheet_number (exact match)"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sheet_number': 'A-101'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['text'], "Note on A-101")

    def test_filter_by_sheet_title_partial(self):
        """Test filtering by sheet_title (case-insensitive partial match)"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sheet_title': 'mechanical'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['text'], "Note on M-201")

    def test_filter_by_sheet_number_is_null(self):
        """Test filtering for records where sheet_number is null"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sheet_number_is_null': 'true'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['text'], "Note without sheet")

    def test_filter_by_sheet_title_is_null(self):
        """Test filtering for records where sheet_title is null"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sheet_title_is_null': 'true'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['text'], "Note without sheet")

    def test_sort_by_sheet_number_asc(self):
        """Test sorting by sheet_number ascending"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'sheet_number', 'sort_direction': 'asc'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        self.assertEqual(len(results), 3)
        # PostgreSQL: nulls sort last in ascending order
        self.assertEqual(results[0]['sheet_number'], 'A-101')
        self.assertEqual(results[1]['sheet_number'], 'M-201')
        self.assertIsNone(results[2]['sheet_number'])

    def test_sort_by_sheet_number_desc(self):
        """Test sorting by sheet_number descending"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'sheet_number', 'sort_direction': 'desc'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        self.assertEqual(len(results), 3)
        # PostgreSQL: nulls sort first in descending order
        self.assertIsNone(results[0]['sheet_number'])
        self.assertEqual(results[1]['sheet_number'], 'M-201')
        self.assertEqual(results[2]['sheet_number'], 'A-101')

    def test_sort_by_sheet_title_asc(self):
        """Test sorting by sheet_title ascending"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'sheet_title', 'sort_direction': 'asc'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        self.assertEqual(len(results), 3)
        # PostgreSQL: nulls sort last in ascending order
        self.assertEqual(results[0]['sheet_title'], 'First Floor Plan')
        self.assertEqual(results[1]['sheet_title'], 'Mechanical Plan')
        self.assertIsNone(results[2]['sheet_title'])

    def test_sort_by_sheet_title_desc(self):
        """Test sorting by sheet_title descending"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sort_column': 'sheet_title', 'sort_direction': 'desc'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results']
        self.assertEqual(len(results), 3)
        # PostgreSQL: nulls sort first in descending order
        self.assertIsNone(results[0]['sheet_title'])
        self.assertEqual(results[1]['sheet_title'], 'Mechanical Plan')
        self.assertEqual(results[2]['sheet_title'], 'First Floor Plan')

    def test_sheet_fields_in_response(self):
        """Test that sheet_number and sheet_title appear in API response"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(len(response.data['results']) > 0)

        # Check that sheet fields are present in the response
        first_result = response.data['results'][0]
        self.assertIn('sheet_number', first_result)
        self.assertIn('sheet_title', first_result)

    def test_all_filter_vals_includes_sheet_numbers(self):
        """Test that all_filter_vals includes sheet_numbers list"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('all_filter_vals', response.data)
        self.assertIn('sheet_numbers', response.data['all_filter_vals'])

        sheet_numbers = response.data['all_filter_vals']['sheet_numbers']
        self.assertIn('A-101', sheet_numbers)
        self.assertIn('M-201', sheet_numbers)
        # Null should not be in the list
        self.assertNotIn(None, sheet_numbers)

    def test_all_filter_vals_includes_sheet_titles(self):
        """Test that all_filter_vals includes sheet_titles list"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('sheet_titles', response.data['all_filter_vals'])

        sheet_titles = response.data['all_filter_vals']['sheet_titles']
        self.assertIn('First Floor Plan', sheet_titles)
        self.assertIn('Mechanical Plan', sheet_titles)
        self.assertNotIn(None, sheet_titles)

    def test_all_filter_vals_has_null_sheet_number_flag(self):
        """Test that has_null_sheet_number flag is correctly set"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('has_null_sheet_number', response.data['all_filter_vals'])
        self.assertTrue(response.data['all_filter_vals']['has_null_sheet_number'])

    def test_all_filter_vals_has_null_sheet_title_flag(self):
        """Test that has_null_sheet_title flag is correctly set"""
        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('has_null_sheet_title', response.data['all_filter_vals'])
        self.assertTrue(response.data['all_filter_vals']['has_null_sheet_title'])

    def test_all_filter_vals_no_nulls_when_all_have_values(self):
        """Test has_null flags are False when all records have values"""
        # Delete the page without sheet values
        self.page3.delete()

        url = reverse('deliverables:drawing-note-list', kwargs={'project_id': self.project.id})
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(response.data['all_filter_vals']['has_null_sheet_number'])
        self.assertFalse(response.data['all_filter_vals']['has_null_sheet_title'])

    def test_export_with_sheet_number_is_null_filter(self):
        """Test export endpoint supports sheet_number_is_null filter"""
        import openpyxl
        from io import BytesIO

        url = reverse('deliverables:drawing-note-export', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sheet_number_is_null': 'true'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        worksheet = workbook.active

        # Should only have header + 1 row (the note without sheet number)
        self.assertEqual(worksheet.max_row, 2)

    def test_export_with_sheet_title_is_null_filter(self):
        """Test export endpoint supports sheet_title_is_null filter"""
        import openpyxl
        from io import BytesIO

        url = reverse('deliverables:drawing-note-export', kwargs={'project_id': self.project.id})
        response = self.client.get(url, {'sheet_title_is_null': 'true'})

        self.assertEqual(response.status_code, status.HTTP_200_OK)

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        worksheet = workbook.active

        # Should only have header + 1 row
        self.assertEqual(worksheet.max_row, 2)
