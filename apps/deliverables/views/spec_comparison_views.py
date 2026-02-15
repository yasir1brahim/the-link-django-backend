import json
import logging
import uuid
import requests
import boto3
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from django.conf import settings
from django.db import transaction, IntegrityError
from django.db.models import Subquery, OuterRef, Count, Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response

from ..models import (
    Project,
    ProjectVersion,
    SpecComparison,
    SpecComparisonStatus,
    SpecComparisonWebhookEvent,
    SpecConflict,
    SkippedNote,
    DrawingNote,
    DrawingExtraction,
    DrawingExtractionStatus,
    SpecSection,
)
from ..serializers.spec_comparison_serializers import (
    SpecComparisonWebhookSerializer,
    TriggerSpecComparisonSerializer,
    TriggerSpecComparisonResponseSerializer,
    SpecComparisonSummarySerializer,
    SpecConflictReadSerializer,
    SkippedNoteReadSerializer,
    SpecComparisonListSerializer,
)
from apps.utils.feature_flags import is_drawing_spec_comparison_active

# Initialize S3 client at module level (matches existing pattern)
s3 = boto3.client(
    "s3",
    region_name=settings.AWS_REGION,
    aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
    aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
)

logger = logging.getLogger(__name__)


def parse_s3_uri_to_key(s3_uri: str) -> str:
    """Extract S3 key from s3://bucket/key URI.

    Only accepts s3:// URIs (validated by serializer).
    """
    # s3://bucket/path/to/file.pdf -> path/to/file.pdf
    parts = s3_uri[5:].split('/', 1)  # Remove 's3://'
    if len(parts) > 1:
        return parts[1]
    raise ValueError(f"Invalid S3 URI format: {s3_uri}")


def safe_int(value: str) -> int | None:
    """Safely convert string to int, return None if invalid."""
    try:
        return int(value)
    except (ValueError, TypeError):
        return None


@api_view(['POST'])
@permission_classes([AllowAny])
def spec_comparison_webhook(request):
    """
    Webhook endpoint for receiving spec comparison results from AWS Lambda.
    Idempotent: duplicate event_ids are ignored.
    Atomic: all changes within a single transaction.
    """
    # Validate payload
    serializer = SpecComparisonWebhookSerializer(data=request.data)
    if not serializer.is_valid():
        return Response(
            {"error": "Invalid payload", "details": serializer.errors},
            status=status.HTTP_400_BAD_REQUEST
        )

    data = serializer.validated_data
    event_id = data['event_id']
    comparison_id = data['comparison_id']
    new_status = data['status']

    with transaction.atomic():
        # Lock comparison row to prevent races
        try:
            comparison = SpecComparison.objects.select_for_update().get(id=comparison_id)
        except SpecComparison.DoesNotExist:
            return Response(
                {"error": "Comparison not found"},
                status=status.HTTP_404_NOT_FOUND
            )

        # Verify event_id matches
        if comparison.event_id != event_id:
            return Response(
                {"error": "Event ID mismatch"},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Create webhook event for idempotency (unique constraint on event_id)
        try:
            SpecComparisonWebhookEvent.objects.create(
                comparison=comparison,
                event_id=event_id,
                new_status=new_status,
            )
        except IntegrityError:
            # Duplicate event_id - already processed
            return Response({"status": "already_processed"}, status=status.HTTP_200_OK)

        # Update comparison metadata
        comparison.status = new_status
        comparison.notes_processed = data['notes_processed']
        comparison.notes_skipped = data['notes_skipped']
        comparison.spec_files_processed = data['spec_files_processed']
        comparison.notes_with_mismatch = data['notes_with_mismatch']
        comparison.completed_at = timezone.now()

        if new_status == SpecComparisonStatus.FAILED:
            comparison.error_message = data.get('error_message')
            comparison.save()
            # FAILED: Do not store conflicts/skipped notes (no partial data)
            return Response({"status": "processed"}, status=status.HTTP_200_OK)

        comparison.save()

        # SUCCESS/PARTIAL_SUCCESS: Store conflicts and skipped notes

        # Build note ID lookup for FK linking (safely handle non-integer IDs)
        conflict_note_ids = [c['note_id'] for c in data['conflicts']]
        skipped_note_ids = [s['note_id'] for s in data['skipped_notes']]
        all_note_ids_raw = set(conflict_note_ids + skipped_note_ids)

        # Safely convert to integers, filtering out invalid IDs
        valid_note_ids = [safe_int(nid) for nid in all_note_ids_raw]
        valid_note_ids = [nid for nid in valid_note_ids if nid is not None]

        notes_by_id = {}
        if valid_note_ids:
            # Scope note lookup to the comparison's project/version for security
            notes_by_id = {
                str(n.id): n
                for n in DrawingNote.objects.filter(
                    id__in=valid_note_ids,
                    section__page__extraction__drawing_file__project=comparison.project,
                    section__page__extraction__drawing_file__project_version=comparison.project_version,
                )
            }

        # Create conflicts (with batch_size for large lists)
        if data['conflicts']:
            conflicts_to_create = [
                SpecConflict(
                    comparison=comparison,
                    note=notes_by_id.get(c['note_id']),
                    note_id_from_lambda=c['note_id'],
                    note_text=c['note_text'],
                    spec_text=c['spec_text'],
                    spec_file_s3_key=parse_s3_uri_to_key(c['spec_source_file']),
                    spec_page_number=c['spec_page_number'],
                    spec_masterformat_number=c['spec_masterformat_number'],
                    confidence=c['confidence'],
                    reason=c['reason'],
                    pdf_locations=c.get('pdf_locations', []),
                )
                for c in data['conflicts']
            ]
            SpecConflict.objects.bulk_create(conflicts_to_create, batch_size=500)

        # Create skipped notes (with batch_size for large lists)
        if data['skipped_notes']:
            skipped_to_create = [
                SkippedNote(
                    comparison=comparison,
                    note=notes_by_id.get(s['note_id']),
                    note_id_from_lambda=s['note_id'],
                    disciplines=s.get('disciplines', []),
                    sheet_discipline=s.get('sheet_discipline'),
                    reason=s['reason'],
                    detail=s.get('detail'),
                )
                for s in data['skipped_notes']
            ]
            SkippedNote.objects.bulk_create(skipped_to_create, batch_size=500)

    return Response({"status": "processed"}, status=status.HTTP_200_OK)


# --- Trigger Endpoint Helper Functions ---


def get_notes_queryset(project, project_version):
    """Get notes from the latest successful extraction per drawing file.

    Includes both SUCCESS and PARTIAL_SUCCESS extractions to match
    the DrawingExtractionStatus enum.
    """
    # Subquery: latest successful extraction ID per drawing file
    latest_extraction_subquery = DrawingExtraction.objects.filter(
        drawing_file=OuterRef('section__page__extraction__drawing_file'),
        drawing_file__project=project,
        drawing_file__project_version=project_version,
        status__in=[DrawingExtractionStatus.SUCCESS, DrawingExtractionStatus.PARTIAL_SUCCESS],
    ).order_by('-created_at').values('id')[:1]

    return DrawingNote.objects.filter(
        section__page__extraction__drawing_file__project=project,
        section__page__extraction__drawing_file__project_version=project_version,
        section__page__extraction__id=Subquery(latest_extraction_subquery),
    ).select_related('section__page')


def get_specs_queryset(project, project_version):
    """Get spec sections with valid S3 keys."""
    return SpecSection.objects.filter(
        document__project=project,
        document__project_version=project_version,
        file_s3_key__isnull=False,
    ).exclude(
        file_s3_key=''
    ).select_related('masterformat_section')


def build_comparison_payload(project, project_version, comparison_id, event_id):
    """Build the payload to send to lambda."""
    drawing_notes = get_notes_queryset(project, project_version)
    spec_sections = get_specs_queryset(project, project_version)

    payload = {
        "comparison_id": comparison_id,
        "event_id": event_id,
        "notes": [
            {
                "id": str(note.id),
                "text": note.text,
                "category": note.category,
                "sheet_number": note.section.page.sheet_number,
                "sheet_title": note.section.page.sheet_title,
                "disciplines": note.disciplines or [],
                "discipline_confidence": note.discipline_confidence,
                "sheet_discipline": note.section.page.sheet_discipline,
                "sheet_discipline_confidence": note.section.page.sheet_discipline_confidence,
            }
            for note in drawing_notes
        ],
        "spec_files": [
            {
                "s3_uri": f"s3://{settings.S3_BUCKET}/{spec.file_s3_key}",
                "s3_key": spec.file_s3_key,
                "masterformat_number": spec.masterformat_section.masterformat_number,
            }
            for spec in spec_sections
        ],
        "callback_url": settings.BACKEND_SPEC_COMPARISON_CALLBACK_URL,
    }

    return payload


def upload_payload_to_s3(payload, comparison_id):
    """Upload payload JSON to S3 and return the key."""
    payload_s3_key = f"spec-comparisons/{comparison_id}/payload.json"
    s3.put_object(
        Bucket=settings.S3_BUCKET,
        Key=payload_s3_key,
        Body=json.dumps(payload),
        ContentType='application/json',
    )
    return payload_s3_key


def invoke_spec_comparison_lambda(payload_s3_key, comparison_id, event_id):
    """Invoke the lambda function with S3 pointer.

    Raises exception if lambda returns non-2xx status or times out unexpectedly.
    ReadTimeout is expected for large payloads as lambda processes asynchronously.
    """
    if not settings.SPEC_COMPARISON_LAMBDA_FUNCTION_URL:
        raise ValueError("SPEC_COMPARISON_LAMBDA_FUNCTION_URL is not set")

    lambda_payload = {
        "payload_s3_uri": f"s3://{settings.S3_BUCKET}/{payload_s3_key}",
        "callback_url": settings.BACKEND_SPEC_COMPARISON_CALLBACK_URL,
        "comparison_id": comparison_id,
        "event_id": event_id,
    }

    try:
        response = requests.post(
            settings.SPEC_COMPARISON_LAMBDA_FUNCTION_URL,
            json=lambda_payload,
            timeout=2
        )
        # Check for non-2xx status codes (lambda invocation errors)
        response.raise_for_status()
    except requests.exceptions.ReadTimeout:
        # If we timed out, the lambda is processing - this is expected
        pass
    except requests.exceptions.HTTPError as e:
        # Lambda returned an error status code
        raise Exception(f"Lambda invocation failed: {e.response.status_code} - {e.response.text[:200]}")


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def trigger_spec_comparison(request, project_id):
    """
    Trigger a new spec comparison for a project.

    POST /api/deliverables/projects/{project_id}/trigger-spec-comparison/
    """
    # Get project
    try:
        project = Project.objects.get(id=project_id)
    except Project.DoesNotExist:
        return Response(
            {"error": "Project not found"},
            status=status.HTTP_404_NOT_FOUND
        )

    # Check project membership
    if not request.user.is_member_of_project(project_id):
        return Response(
            {"error": "Not authorized to access this project"},
            status=status.HTTP_403_FORBIDDEN
        )

    # Check feature flag
    if not is_drawing_spec_comparison_active(request.user, project.team, project):
        return Response(
            {"error": "Feature not enabled for this project"},
            status=status.HTTP_403_FORBIDDEN
        )

    # Parse request body
    serializer = TriggerSpecComparisonSerializer(data=request.data)
    if not serializer.is_valid():
        return Response(
            {"error": "Invalid request", "details": serializer.errors},
            status=status.HTTP_400_BAD_REQUEST
        )

    # Resolve project version (required parameter)
    project_version_id = serializer.validated_data['project_version_id']
    try:
        project_version = ProjectVersion.objects.get(
            id=project_version_id,
            project=project
        )
    except ProjectVersion.DoesNotExist:
        return Response(
            {"error": "Project version not found or does not belong to this project"},
            status=status.HTTP_400_BAD_REQUEST
        )

    # Pre-validate: count notes and specs
    notes_count = get_notes_queryset(project, project_version).count()
    specs_count = get_specs_queryset(project, project_version).count()

    if notes_count == 0:
        return Response(
            {"error": "No drawing notes found for comparison"},
            status=status.HTTP_400_BAD_REQUEST
        )
    if specs_count == 0:
        return Response(
            {"error": "No spec sections found for comparison"},
            status=status.HTTP_400_BAD_REQUEST
        )

    # Create comparison with server-generated event_id
    # Set PROCESSING status BEFORE invoking lambda to prevent race condition
    # where webhook returns before we update status
    event_id = str(uuid.uuid4())
    comparison = SpecComparison.objects.create(
        project=project,
        project_version=project_version,
        triggered_by=request.user,
        status=SpecComparisonStatus.PROCESSING,  # Start as PROCESSING, not PENDING
        event_id=event_id,
        started_at=timezone.now(),
    )

    try:
        # Build and upload payload
        payload = build_comparison_payload(project, project_version, comparison.id, event_id)
        payload_s3_key = upload_payload_to_s3(payload, comparison.id)
        comparison.payload_s3_key = payload_s3_key
        comparison.save()

        # Invoke lambda (status already PROCESSING)
        invoke_spec_comparison_lambda(payload_s3_key, comparison.id, event_id)

    except Exception as e:
        comparison.status = SpecComparisonStatus.FAILED
        comparison.error_message = str(e)
        comparison.completed_at = timezone.now()
        comparison.save()
        logger.error(f"Failed to trigger spec comparison {comparison.id}: {e}")

    response_serializer = TriggerSpecComparisonResponseSerializer(comparison)
    return Response(response_serializer.data, status=status.HTTP_201_CREATED)


# --- Read Endpoints ---


# Sortable columns mapping for spec conflicts
SORTABLE_COLUMNS = {
    'sheet_number': 'note__section__page__sheet_number',
    'note_text': 'note_text',
    'spec_text': 'spec_text',
    'spec_masterformat_number': 'spec_masterformat_number',
    'reason': 'reason',
}


class SpecComparisonPagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = 'limit'  # Use 'limit' for consistency with DrawingNotePagination
    max_page_size = 100


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_spec_conflicts(request, project_id):
    """
    Get spec conflicts from the latest successful comparison.

    GET /api/deliverables/projects/{project_id}/spec-conflicts/
    Query params:
        - comparison_id: Filter to specific comparison
        - project_version_id: Required - the version to query
        - page, limit: Pagination
    """
    # Get project and check access
    try:
        project = Project.objects.get(id=project_id)
    except Project.DoesNotExist:
        return Response(
            {"error": "Project not found"},
            status=status.HTTP_404_NOT_FOUND
        )

    if not request.user.is_member_of_project(project_id):
        return Response(
            {"error": "Not authorized to access this project"},
            status=status.HTTP_403_FORBIDDEN
        )

    # Get version filter (required)
    project_version_id = request.query_params.get('project_version_id')
    if not project_version_id:
        return Response(
            {"error": "project_version_id is required"},
            status=status.HTTP_400_BAD_REQUEST
        )
    try:
        project_version = ProjectVersion.objects.get(
            id=int(project_version_id),
            project=project
        )
    except (ValueError, ProjectVersion.DoesNotExist):
        return Response(
            {"error": "Invalid project_version_id"},
            status=status.HTTP_400_BAD_REQUEST
        )

    # Get comparison
    comparison_id = request.query_params.get('comparison_id')
    if comparison_id:
        try:
            comparison_id_int = int(comparison_id)
        except (ValueError, TypeError):
            return Response(
                {"error": "comparison_id must be an integer"},
                status=status.HTTP_400_BAD_REQUEST
            )
        try:
            comparison = SpecComparison.objects.get(
                id=comparison_id_int,
                project=project,
                project_version=project_version
            )
        except SpecComparison.DoesNotExist:
            return Response(
                {"error": "Comparison not found or does not belong to this project/version"},
                status=status.HTTP_404_NOT_FOUND
            )
    else:
        # Get latest successful comparison for this version
        comparison = SpecComparison.objects.filter(
            project=project,
            project_version=project_version,
            status__in=[SpecComparisonStatus.SUCCESS, SpecComparisonStatus.PARTIAL_SUCCESS],
        ).order_by('-completed_at').first()

    if not comparison:
        return Response({
            'count': 0,
            'next': None,
            'previous': None,
            'comparison': None,
            'filter_options': {
                'sheet_numbers': [],
                'spec_masterformat_numbers': [],
                'reasons': [],
            },
            'results': [],
        })

    # Get filter options from unfiltered queryset (for dropdown population)
    all_conflicts = SpecConflict.objects.filter(comparison=comparison).select_related(
        'note__section__page'
    )

    filter_options = {
        'sheet_numbers': sorted(list(
            all_conflicts.exclude(note__section__page__sheet_number__isnull=True)
            .values_list('note__section__page__sheet_number', flat=True)
            .distinct()
        )),
        'spec_masterformat_numbers': sorted(list(
            all_conflicts.values_list('spec_masterformat_number', flat=True).distinct()
        )),
        'reasons': sorted(list(
            all_conflicts.values_list('reason', flat=True).distinct()
        )),
    }

    # Get conflicts with pagination (select_related to avoid N+1 on note access)
    conflicts = SpecConflict.objects.filter(comparison=comparison).select_related(
        'note__section__page__drawing_file'
    )

    # Apply filters
    search = request.query_params.get('search')
    if search:
        conflicts = conflicts.filter(
            Q(note_text__icontains=search) |
            Q(spec_text__icontains=search) |
            Q(reason__icontains=search)
        )

    sheet_number_filter = request.query_params.get('sheet_number')
    if sheet_number_filter:
        conflicts = conflicts.filter(note__section__page__sheet_number=sheet_number_filter)

    spec_masterformat_filter = request.query_params.get('spec_masterformat_number')
    if spec_masterformat_filter:
        conflicts = conflicts.filter(spec_masterformat_number=spec_masterformat_filter)

    reason_filter = request.query_params.get('reason')
    if reason_filter:
        conflicts = conflicts.filter(reason=reason_filter)

    # Apply sorting
    sort_column = request.query_params.get('sort_column')
    sort_direction = request.query_params.get('sort_direction', 'asc')

    if sort_column and sort_column in SORTABLE_COLUMNS:
        order_field = SORTABLE_COLUMNS[sort_column]
        if sort_direction == 'desc':
            order_field = f'-{order_field}'
        conflicts = conflicts.order_by(order_field, 'id')  # Secondary sort by id for stability
    else:
        conflicts = conflicts.order_by('id')

    paginator = SpecComparisonPagination()
    page = paginator.paginate_queryset(conflicts, request)

    serializer = SpecConflictReadSerializer(page, many=True, context={'request': request})

    response = paginator.get_paginated_response(serializer.data)
    response.data['comparison'] = SpecComparisonSummarySerializer(comparison).data
    response.data['filter_options'] = filter_options

    return response


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_skipped_notes(request, project_id):
    """
    Get skipped notes from the latest successful comparison.

    GET /api/deliverables/projects/{project_id}/skipped-notes/
    Query params:
        - comparison_id: Filter to specific comparison
        - project_version_id: Required - the version to query
        - reason: Filter by skip reason
        - page, limit: Pagination
    """
    # Get project and check access
    try:
        project = Project.objects.get(id=project_id)
    except Project.DoesNotExist:
        return Response(
            {"error": "Project not found"},
            status=status.HTTP_404_NOT_FOUND
        )

    if not request.user.is_member_of_project(project_id):
        return Response(
            {"error": "Not authorized to access this project"},
            status=status.HTTP_403_FORBIDDEN
        )

    # Get version filter (required)
    project_version_id = request.query_params.get('project_version_id')
    if not project_version_id:
        return Response(
            {"error": "project_version_id is required"},
            status=status.HTTP_400_BAD_REQUEST
        )
    try:
        project_version = ProjectVersion.objects.get(
            id=int(project_version_id),
            project=project
        )
    except (ValueError, ProjectVersion.DoesNotExist):
        return Response(
            {"error": "Invalid project_version_id"},
            status=status.HTTP_400_BAD_REQUEST
        )

    # Get comparison
    comparison_id = request.query_params.get('comparison_id')
    if comparison_id:
        try:
            comparison_id_int = int(comparison_id)
        except (ValueError, TypeError):
            return Response(
                {"error": "comparison_id must be an integer"},
                status=status.HTTP_400_BAD_REQUEST
            )
        try:
            comparison = SpecComparison.objects.get(
                id=comparison_id_int,
                project=project,
                project_version=project_version
            )
        except SpecComparison.DoesNotExist:
            return Response(
                {"error": "Comparison not found or does not belong to this project/version"},
                status=status.HTTP_404_NOT_FOUND
            )
    else:
        # Get latest successful comparison for this version
        comparison = SpecComparison.objects.filter(
            project=project,
            project_version=project_version,
            status__in=[SpecComparisonStatus.SUCCESS, SpecComparisonStatus.PARTIAL_SUCCESS],
        ).order_by('-completed_at').first()

    if not comparison:
        return Response({
            'count': 0,
            'next': None,
            'previous': None,
            'results': [],
        })

    # Get skipped notes with optional reason filter (select_related to avoid N+1 on note access)
    skipped_notes = SkippedNote.objects.filter(comparison=comparison).select_related('note').order_by('id')

    reason_filter = request.query_params.get('reason')
    if reason_filter:
        skipped_notes = skipped_notes.filter(reason=reason_filter)

    paginator = SpecComparisonPagination()
    page = paginator.paginate_queryset(skipped_notes, request)

    serializer = SkippedNoteReadSerializer(page, many=True)

    return paginator.get_paginated_response(serializer.data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def list_spec_comparisons(request, project_id):
    """
    List all spec comparisons for a project.

    GET /api/deliverables/projects/{project_id}/spec-comparisons/
    """
    # Get project and check access
    try:
        project = Project.objects.get(id=project_id)
    except Project.DoesNotExist:
        return Response(
            {"error": "Project not found"},
            status=status.HTTP_404_NOT_FOUND
        )

    if not request.user.is_member_of_project(project_id):
        return Response(
            {"error": "Not authorized to access this project"},
            status=status.HTTP_403_FORBIDDEN
        )

    comparisons = SpecComparison.objects.filter(
        project=project
    ).order_by('-created_at').annotate(
        conflict_count_annotated=Count('conflicts')
    ).select_related('triggered_by')

    paginator = SpecComparisonPagination()
    page = paginator.paginate_queryset(comparisons, request)

    serializer = SpecComparisonListSerializer(page, many=True)

    return paginator.get_paginated_response(serializer.data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_spec_conflicts(request, project_id):
    """
    Export spec conflicts to Excel file.

    GET /api/deliverables/projects/{project_id}/spec-conflicts/export/
    Query params: Same as get_spec_conflicts (filters apply)
    """
    # Get project and check access
    try:
        project = Project.objects.get(id=project_id)
    except Project.DoesNotExist:
        return Response(
            {"error": "Project not found"},
            status=status.HTTP_404_NOT_FOUND
        )

    if not request.user.is_member_of_project(project_id):
        return Response(
            {"error": "Not authorized to access this project"},
            status=status.HTTP_403_FORBIDDEN
        )

    # Get version filter (required)
    project_version_id = request.query_params.get('project_version_id')
    if not project_version_id:
        return Response(
            {"error": "project_version_id is required"},
            status=status.HTTP_400_BAD_REQUEST
        )
    try:
        project_version = ProjectVersion.objects.get(
            id=int(project_version_id),
            project=project
        )
    except (ValueError, ProjectVersion.DoesNotExist):
        return Response(
            {"error": "Invalid project_version_id"},
            status=status.HTTP_400_BAD_REQUEST
        )

    # Get comparison
    comparison_id = request.query_params.get('comparison_id')
    if comparison_id:
        try:
            comparison = SpecComparison.objects.get(
                id=int(comparison_id),
                project=project,
                project_version=project_version
            )
        except (ValueError, SpecComparison.DoesNotExist):
            return Response(
                {"error": "Comparison not found"},
                status=status.HTTP_404_NOT_FOUND
            )
    else:
        comparison = SpecComparison.objects.filter(
            project=project,
            project_version=project_version,
            status__in=[SpecComparisonStatus.SUCCESS, SpecComparisonStatus.PARTIAL_SUCCESS],
        ).order_by('-completed_at').first()

    if not comparison:
        # Return empty Excel file
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Spec Conflicts"
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=spec_conflicts.xlsx'
        workbook.save(response)
        return response

    # Get conflicts with filters
    conflicts = SpecConflict.objects.filter(comparison=comparison).select_related(
        'note__section__page__drawing_file'
    )

    # Apply filters (same as get_spec_conflicts)
    search = request.query_params.get('search')
    if search:
        conflicts = conflicts.filter(
            Q(note_text__icontains=search) |
            Q(spec_text__icontains=search) |
            Q(reason__icontains=search)
        )

    sheet_number_filter = request.query_params.get('sheet_number')
    if sheet_number_filter:
        conflicts = conflicts.filter(note__section__page__sheet_number=sheet_number_filter)

    spec_masterformat_filter = request.query_params.get('spec_masterformat_number')
    if spec_masterformat_filter:
        conflicts = conflicts.filter(spec_masterformat_number=spec_masterformat_filter)

    reason_filter = request.query_params.get('reason')
    if reason_filter:
        conflicts = conflicts.filter(reason=reason_filter)

    conflicts = conflicts.order_by('id')

    # Create workbook
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Spec Conflicts"

    # Define styles
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='202a44', end_color='202a44', fill_type='solid')
    header_alignment = Alignment(wrap_text=True, vertical='center')
    text_alignment = Alignment(wrap_text=True, vertical='center')

    # Headers
    headers = [
        'Drawing #', 'Sheet Title', 'Drawing Content', 'Related Spec Content',
        'Spec Section', 'Spec Page', 'Reason', 'Confidence'
    ]
    for col, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Data rows
    for row_idx, conflict in enumerate(conflicts, start=2):
        # Get drawing info
        sheet_number = None
        sheet_title = None
        if conflict.note and conflict.note.section and conflict.note.section.page:
            page = conflict.note.section.page
            sheet_number = page.sheet_number
            sheet_title = page.sheet_title

        row_data = [
            sheet_number,
            sheet_title,
            conflict.note_text,
            conflict.spec_text,
            conflict.spec_masterformat_number,
            conflict.spec_page_number,
            conflict.reason,
            f"{int(conflict.confidence * 100)}%",
        ]

        for col, value in enumerate(row_data, start=1):
            cell = worksheet.cell(row=row_idx, column=col, value=value)
            cell.alignment = text_alignment

    # Set column widths
    column_widths = [15, 30, 50, 50, 15, 12, 40, 12]
    for col, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=spec_conflicts.xlsx'
    workbook.save(response)

    return response
