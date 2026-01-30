import re

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from django.db import transaction, IntegrityError
from django.db.models import OuterRef, Subquery
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.decorators import api_view, permission_classes, action
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response
from rest_framework import status, viewsets
from rest_framework.pagination import PageNumberPagination

from .main_views import clean_excel_content

from ..models import (
    DrawingExtraction,
    DrawingExtractionStatus,
    DrawingExtractionWebhookEvent,
    DrawingPage,
    DrawingPageExtractionStatus,
    DrawingNoteSection,
    DrawingNote,
    DrawingFile,
)
from ..serializers.drawing_serializers import DrawingNoteReadSerializer
from ..permissions import DrawingNoteAccessPermissions


@api_view(['POST'])
@permission_classes([AllowAny])
def drawing_extraction_webhook(request):
    """
    Webhook endpoint for receiving drawing extraction results from AWS Lambda.

    Idempotent: duplicate event_ids are ignored.
    Atomic: all changes within a single transaction.
    """
    payload = request.data
    event_id = payload.get('event_id')
    extraction_id = payload.get('extraction_id')
    new_status_str = payload.get('new_status')

    if not all([event_id, extraction_id, new_status_str]):
        return Response(
            {"error": "Missing required fields: event_id, extraction_id, new_status"},
            status=status.HTTP_400_BAD_REQUEST
        )

    try:
        new_status = DrawingExtractionStatus(new_status_str)
    except ValueError:
        return Response(
            {"error": f"Invalid status: {new_status_str}"},
            status=status.HTTP_400_BAD_REQUEST
        )

    with transaction.atomic():
        try:
            extraction = DrawingExtraction.objects.select_for_update().get(id=extraction_id)
        except DrawingExtraction.DoesNotExist:
            return Response(
                {"error": f"Extraction {extraction_id} not found"},
                status=status.HTTP_404_NOT_FOUND
            )

        # Try to record the webhook event for idempotency
        try:
            DrawingExtractionWebhookEvent.objects.create(
                extraction=extraction,
                event_id=event_id,
                new_status=new_status.value,
                output_s3_key=payload.get("output_s3_key"),
                payload=payload,
            )
        except IntegrityError:
            # Duplicate event_id - already processed
            return Response(status=status.HTTP_200_OK)

        # Status transition validation
        allowed_next_statuses = {
            DrawingExtractionStatus.PENDING: {
                DrawingExtractionStatus.PROCESSING,
                DrawingExtractionStatus.FAILED
            },
            DrawingExtractionStatus.PROCESSING: {
                DrawingExtractionStatus.SUCCESS,
                DrawingExtractionStatus.PARTIAL_SUCCESS,
                DrawingExtractionStatus.FAILED,
            },
            DrawingExtractionStatus.SUCCESS: set(),
            DrawingExtractionStatus.PARTIAL_SUCCESS: set(),
            DrawingExtractionStatus.FAILED: set(),
        }

        if new_status not in allowed_next_statuses.get(extraction.status, set()):
            # Ignore invalid transitions (e.g., PROCESSING after SUCCESS)
            return Response(status=status.HTTP_200_OK)

        # Process based on new status
        if new_status == DrawingExtractionStatus.PROCESSING:
            extraction.status = DrawingExtractionStatus.PROCESSING
            extraction.started_at = timezone.now()
            extraction.save()

        elif new_status in {DrawingExtractionStatus.SUCCESS, DrawingExtractionStatus.PARTIAL_SUCCESS}:
            extraction.status = new_status
            extraction.completed_at = timezone.now()
            extraction.model_version = payload.get('model_version')
            extraction.processing_time_ms = payload.get('processing_time_ms')
            extraction.output_s3_key = payload.get('output_s3_key')
            extraction.failure_summary = payload.get('failure_summary')
            extraction.save()

            # Update drawing file total_pages
            data = payload.get('data', {})
            if data.get('total_pages'):
                extraction.drawing_file.total_pages = data['total_pages']
                extraction.drawing_file.save()

            # Clear existing pages and recreate (idempotent)
            DrawingPage.objects.filter(extraction=extraction).delete()
            _create_drawing_records(extraction, data)

        elif new_status == DrawingExtractionStatus.FAILED:
            extraction.status = DrawingExtractionStatus.FAILED
            extraction.completed_at = timezone.now()
            extraction.error_message = payload.get('error_message')
            extraction.failure_summary = payload.get('failure_summary')
            extraction.save()

    return Response(status=status.HTTP_200_OK)


def _create_drawing_records(extraction, data):
    """
    Bulk create DrawingPage, DrawingNoteSection, and DrawingNote records.
    Uses bulk_create for performance on large PDFs.
    """
    pages_data = data.get('pages', [])

    def _normalize_page_extraction_status(status_value):
        if status_value is None:
            return status_value

        normalized = str(status_value).lower()
        if normalized == 'error':
            return DrawingPageExtractionStatus.FAILED
        if normalized == 'partial':
            return DrawingPageExtractionStatus.FAILED
        if normalized == 'success':
            return DrawingPageExtractionStatus.SUCCESS
        if normalized == 'no_notes_found':
            return DrawingPageExtractionStatus.NO_NOTES_FOUND
        if normalized == 'failed':
            return DrawingPageExtractionStatus.FAILED

        return status_value

    # First pass: create all pages
    page_objects = []
    for page_data in pages_data:
        page_objects.append(DrawingPage(
            drawing_file=extraction.drawing_file,
            extraction=extraction,
            page_number=page_data['page_number'],
            page_type=page_data['page_type'],
            rotation=page_data.get('rotation', 0),
            rotated_width=page_data.get('rotated_page_width'),
            rotated_height=page_data.get('rotated_page_height'),
            unrotated_width=page_data.get('unrotated_page_width'),
            unrotated_height=page_data.get('unrotated_page_height'),
            extraction_status=_normalize_page_extraction_status(page_data['extraction_status']),
            spec_content=page_data.get('spec_content'),
            sheet_number=page_data.get('sheet_number'),
            sheet_title=page_data.get('sheet_title'),
            sheet_discipline=page_data.get('sheet_discipline'),
            sheet_discipline_confidence=page_data.get('sheet_discipline_confidence'),
        ))

    created_pages = DrawingPage.objects.bulk_create(page_objects)

    # Build page_number -> page mapping
    page_map = {p.page_number: p for p in created_pages}

    # Second pass: create all note sections
    section_objects = []
    section_notes_map = []  # Track (section_index, notes_data) for later

    for page_data in pages_data:
        page = page_map[page_data['page_number']]
        for section_data in page_data.get('note_sections', []):
            section_index = len(section_objects)
            
            rotated_header_bbox = section_data.get('rotated_header_bbox')
            unrotated_header_bbox = section_data.get('unrotated_header_bbox')
            # Fallback for older payloads or consistency
            header_bbox = unrotated_header_bbox or rotated_header_bbox

            section_objects.append(DrawingNoteSection(
                page=page,
                header=section_data['header'],
                header_bbox=header_bbox,
                rotated_header_bbox=rotated_header_bbox,
                unrotated_header_bbox=unrotated_header_bbox,
            ))
            section_notes_map.append((section_index, section_data.get('notes', [])))

    created_sections = DrawingNoteSection.objects.bulk_create(section_objects)

    # Third pass: create all notes
    note_objects = []
    for section_index, notes_data in section_notes_map:
        section = created_sections[section_index]
        for note_data in notes_data:
            note_objects.append(DrawingNote(
                section=section,
                note_number=note_data['note_number'],
                category=note_data['category'],
                text=note_data['text'],
                bounding_box=note_data.get('unrotated_bounding_box') or note_data.get('bounding_box'),
                raw_bounding_box=note_data.get('rotated_bounding_box') or note_data.get('raw_bounding_box') or note_data.get('bounding_box'),
                rotated_bounding_box=note_data.get('rotated_bounding_box'),
                unrotated_bounding_box=note_data.get('unrotated_bounding_box'),
                source_blocks=note_data.get('source_blocks'),
                drawing_references=note_data.get('drawing_references'),
                disciplines=note_data.get('disciplines', []),
                discipline_confidence=note_data.get('discipline_confidence'),
            ))

    DrawingNote.objects.bulk_create(note_objects)


class DrawingNotePagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = 'limit'
    max_page_size = 100


class DrawingNoteViewSet(viewsets.ReadOnlyModelViewSet):
    """
    Read-only ViewSet for listing drawing notes.
    Mirrors SubmittalItemViewSet pattern with filtering and pagination.
    """
    permission_classes = [IsAuthenticated, DrawingNoteAccessPermissions]
    pagination_class = DrawingNotePagination
    serializer_class = DrawingNoteReadSerializer

    # Sorting configuration
    allowed_sort_columns = ['drawing_file_name', 'category', 'text', 'sheet_number', 'sheet_title']
    allowed_sort_directions = ['asc', 'desc']
    sort_column_mapping = {
        'drawing_file_name': 'section__page__drawing_file__file_name',
        'category': 'category',
        'text': 'text',
        'sheet_number': 'section__page__sheet_number',
        'sheet_title': 'section__page__sheet_title',
    }

    def apply_sorting(self, queryset):
        """
        Apply sorting to the queryset based on query parameters.
        Returns queryset with default ordering if no valid sort_column provided.
        """
        sort_column = self.request.query_params.get('sort_column')
        sort_direction = self.request.query_params.get('sort_direction', 'asc')

        if sort_column:
            # Validate sort_column
            if sort_column not in self.allowed_sort_columns:
                # Invalid sort_column - fall back to default ordering
                return queryset.order_by(
                    'section__page__drawing_file__file_name',
                    'section__page__page_number',
                    'section__header',
                    'note_number',
                )

            # Validate sort_direction
            if sort_direction not in self.allowed_sort_directions:
                sort_direction = 'asc'

            # Map API column name to database field
            db_field = self.sort_column_mapping.get(sort_column, sort_column)

            # Apply descending prefix if needed
            if sort_direction == 'desc':
                db_field = f'-{db_field}'

            return queryset.order_by(db_field)

        # Default ordering when no sort_column specified
        return queryset.order_by(
            'section__page__drawing_file__file_name',
            'section__page__page_number',
            'section__header',
            'note_number',
        )

    def get_queryset(self):
        project_id = self.kwargs['project_id']
        queryset = DrawingNote.objects.filter(
            section__page__drawing_file__project_id=project_id
        ).select_related(
            'section__page__drawing_file',
            'section__page__extraction',
        )

        # Filter by project_version_id
        version_id = self.request.query_params.get('project_version_id')
        if version_id:
            queryset = queryset.filter(
                section__page__drawing_file__project_version_id=version_id
            )

        # Filter by category
        category = self.request.query_params.get('category')
        if category:
            queryset = queryset.filter(category=category)

        # Filter by drawing_file_id
        drawing_file_id = self.request.query_params.get('drawing_file_id')
        if drawing_file_id:
            queryset = queryset.filter(
                section__page__drawing_file_id=drawing_file_id
            )

        # Filter by sheet_number
        sheet_number = self.request.query_params.get('sheet_number')
        if sheet_number:
            queryset = queryset.filter(section__page__sheet_number=sheet_number)

        # Filter by sheet_number_is_null
        sheet_number_is_null = self.request.query_params.get('sheet_number_is_null')
        if sheet_number_is_null and sheet_number_is_null.lower() == 'true':
            queryset = queryset.filter(section__page__sheet_number__isnull=True)

        # Filter by sheet_title
        sheet_title = self.request.query_params.get('sheet_title')
        if sheet_title:
            queryset = queryset.filter(section__page__sheet_title__icontains=sheet_title)

        # Filter by sheet_title_is_null
        sheet_title_is_null = self.request.query_params.get('sheet_title_is_null')
        if sheet_title_is_null and sheet_title_is_null.lower() == 'true':
            queryset = queryset.filter(section__page__sheet_title__isnull=True)

        # Search in text (case-insensitive substring)
        search = self.request.query_params.get('search')
        if search:
            queryset = queryset.filter(text__icontains=search)

        # Apply sorting based on query parameters
        queryset = self.apply_sorting(queryset)

        return queryset

    def _get_processing_status(self, project_id, version_id=None):
        """
        Get processing status for all drawing files in the project/version.
        Returns status object with file counts and details for files still processing.
        """
        # Get all drawing files for this project/version
        files_qs = DrawingFile.objects.filter(project_id=project_id)
        if version_id:
            files_qs = files_qs.filter(project_version_id=version_id)

        # Annotate with latest extraction status
        latest_extraction = DrawingExtraction.objects.filter(
            drawing_file=OuterRef('pk')
        ).order_by('-created_at')

        files_with_status = files_qs.annotate(
            latest_status=Subquery(latest_extraction.values('status')[:1])
        )

        # Count by status
        files_processing = 0
        files_completed = 0
        files_failed = 0
        processing_files = []

        for f in files_with_status:
            file_status = f.latest_status
            if file_status in [DrawingExtractionStatus.PENDING, DrawingExtractionStatus.PROCESSING]:
                files_processing += 1
                processing_files.append({
                    'id': f.id,
                    'name': f.file_name,
                    'status': file_status or DrawingExtractionStatus.PENDING,
                })
            elif file_status in [DrawingExtractionStatus.SUCCESS, DrawingExtractionStatus.PARTIAL_SUCCESS]:
                files_completed += 1
            elif file_status == DrawingExtractionStatus.FAILED:
                files_failed += 1
            elif file_status is None:
                # No extraction yet - treat as pending
                files_processing += 1
                processing_files.append({
                    'id': f.id,
                    'name': f.file_name,
                    'status': DrawingExtractionStatus.PENDING,
                })

        return {
            'is_processing': files_processing > 0,
            'files_processing': files_processing,
            'files_completed': files_completed,
            'files_failed': files_failed,
            'files': processing_files,
        }

    def list(self, request, *args, **kwargs):
        project_id = self.kwargs['project_id']
        version_id = request.query_params.get('project_version_id')

        queryset = self.filter_queryset(self.get_queryset())

        # Get filter values for UI dropdowns
        # Get unique drawing files with id and name
        drawing_files_qs = DrawingFile.objects.filter(
            pages__note_sections__notes__in=queryset
        ).distinct().values('id', 'file_name')

        # Get unique sheet numbers and titles from the queryset
        sheet_numbers_raw = list(queryset.values_list(
            'section__page__sheet_number', flat=True
        ).distinct())
        sheet_titles_raw = list(queryset.values_list(
            'section__page__sheet_title', flat=True
        ).distinct())

        # Separate nulls from values for the filter lists
        sheet_numbers = sorted([sn for sn in sheet_numbers_raw if sn is not None])
        sheet_titles = sorted([st for st in sheet_titles_raw if st is not None])
        has_null_sheet_number = None in sheet_numbers_raw
        has_null_sheet_title = None in sheet_titles_raw

        all_filter_vals = {
            'category': list(set(queryset.values_list('category', flat=True))),
            'drawing_files': [{'id': df['id'], 'name': df['file_name']} for df in drawing_files_qs],
            'sheet_numbers': sheet_numbers,
            'sheet_titles': sheet_titles,
            'has_null_sheet_number': has_null_sheet_number,
            'has_null_sheet_title': has_null_sheet_title,
        }

        # Get processing status
        processing_status = self._get_processing_status(project_id, version_id)

        page = self.paginate_queryset(queryset)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            response = self.get_paginated_response(serializer.data)
            response.data['all_filter_vals'] = all_filter_vals
            response.data['total_count'] = queryset.count()
            response.data['processing_status'] = processing_status
            return response

        serializer = self.get_serializer(queryset, many=True)
        return Response({
            'results': serializer.data,
            'all_filter_vals': all_filter_vals,
            'total_count': queryset.count(),
            'processing_status': processing_status,
        })

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request, *args, **kwargs):
        """
        Export drawing notes to XLSX format.
        Supports the same filters as the list endpoint.
        """
        queryset = self.filter_queryset(self.get_queryset())

        # Create a workbook and select the active worksheet
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Drawing Notes"

        # Define the styles (matching submittal export pattern)
        text_alignment = Alignment(wrap_text=True, vertical='center')
        header_alignment = Alignment(wrap_text=True, vertical='center')
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='202a44', end_color='202a44', fill_type='solid')
        header_border = Border(right=Side(border_style='thin', color='FFFFFF'))

        # Define the headers
        headers = ['File Name', 'Sheet #', 'Sheet Title', 'Page', 'Section', 'Note #', 'Category', 'Note Text']
        worksheet.append(headers)

        # Style the header row
        for col in range(1, len(headers) + 1):
            cell = worksheet.cell(row=1, column=col)
            cell.alignment = header_alignment
            cell.font = header_font
            cell.fill = header_fill
            cell.border = header_border

        # Write data to the worksheet
        row_idx = 2
        for note in queryset:
            row = [
                note.section.page.drawing_file.file_name,
                note.section.page.sheet_number or '',
                note.section.page.sheet_title or '',
                note.section.page.page_number,
                note.section.header,
                note.note_number,
                note.category,
                clean_excel_content(re.sub(ILLEGAL_CHARACTERS_RE, '', note.text or '')),
            ]
            worksheet.append(row)
            for col in range(1, len(row) + 1):
                cell = worksheet.cell(row=row_idx, column=col)
                cell.alignment = text_alignment
            row_idx += 1

        # Adjust column widths
        column_widths = {
            'A': 40,  # File Name
            'B': 15,  # Sheet #
            'C': 40,  # Sheet Title
            'D': 10,  # Page
            'E': 30,  # Section
            'F': 10,  # Note #
            'G': 20,  # Category
            'H': 100,  # Note Text
        }
        for col_letter, width in column_widths.items():
            worksheet.column_dimensions[col_letter].width = width

        # Create a response object and set the appropriate headers
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=drawing_notes.xlsx'

        # Save the workbook to the response
        workbook.save(response)

        return response
